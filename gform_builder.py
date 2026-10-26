"""Generic tabular-data importer and Google Forms Apps Script generator."""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - handled with a user-facing error
    load_workbook = None


MAX_IMPORT_ROWS = 10_000
MAX_IMPORT_COLUMNS = 100


class GFormImportError(ValueError):
    """Raised when a tabular upload cannot be interpreted safely."""


FIELD_ALIASES = {
    "order": {
        "order", "urutan", "nomorurut", "no", "nomor", "sequence", "index",
    },
    "row_type": {
        "rowtype", "type", "tipe", "jenis", "questiontype", "tipepertanyaan",
        "jenissoal", "itemtype",
    },
    "title": {
        "question", "pertanyaan", "soal", "title", "judul", "titleorquestion",
        "pernyataan", "item", "questiontext", "tekspertanyaan",
    },
    "description": {
        "description", "deskripsi", "helptext", "bantuan", "petunjuk",
        "keterangan", "subtitle", "questiondescription",
    },
    "options": {
        "options", "option", "choices", "choice", "pilihan", "opsi",
        "jawaban", "answeroptions", "choicevalues",
    },
    "required": {
        "required", "wajib", "mandatory", "harusdiisi", "isrequired",
    },
    "section": {
        "section", "sectiontitle", "bagian", "bagianjudul", "pagetitle",
        "halaman", "kategori", "group", "kelompok",
    },
    "section_description": {
        "sectiondescription", "pagedescription", "deskripsibagian",
        "deskripsihalaman", "sectionhelp", "pagehelp",
    },
    "scale_min": {
        "scalemin", "min", "minimum", "nilaimin", "batasbawah",
    },
    "scale_max": {
        "scalemax", "max", "maximum", "nilaimax", "batasatas",
    },
    "scale_min_label": {
        "scaleminlabel", "minlabel", "labelmin", "labelbawah",
    },
    "scale_max_label": {
        "scalemaxlabel", "maxlabel", "labelmax", "labelatas",
    },
    "grid_rows": {
        "rows", "gridrows", "barisgrid", "baris", "statements",
    },
    "grid_columns": {
        "columns", "gridcolumns", "kolomgrid", "kolom", "ratings",
    },
}


FIELD_LABELS = {
    "row_type": "Tipe pertanyaan",
    "title": "Judul/pertanyaan",
    "description": "Deskripsi/bantuan",
    "options": "Pilihan jawaban",
    "required": "Wajib diisi",
    "section": "Nama bagian",
    "section_description": "Deskripsi bagian",
    "order": "Urutan",
    "scale_min": "Skala minimum",
    "scale_max": "Skala maksimum",
    "scale_min_label": "Label minimum",
    "scale_max_label": "Label maksimum",
    "grid_rows": "Baris grid",
    "grid_columns": "Kolom grid",
}


TYPE_ALIASES = {
    "section": {
        "section", "page", "bagian", "halaman", "pagebreak", "sectionbreak",
    },
    "header": {
        "header", "itemheader", "sectionheader", "judulitem", "keterangan",
    },
    "short_text": {
        "shortanswer", "shorttext", "text", "jawabansingkat", "singkat",
        "isian", "textitem",
    },
    "paragraph": {
        "paragraph", "longanswer", "longtext", "paragraf", "uraian", "esai",
    },
    "multiple_choice": {
        "multiplechoice", "pilihanganda", "radio", "singlechoice", "mcq",
    },
    "checkbox": {
        "checkbox", "checkboxes", "kotakcentang", "multipleanswer",
        "multiselect",
    },
    "dropdown": {
        "dropdown", "list", "daftarpilihan", "select",
    },
    "linear_scale": {
        "linearscale", "scale", "skala", "skalalinear", "likert",
    },
    "date": {"date", "tanggal"},
    "time": {"time", "waktu", "jam"},
    "grid": {
        "grid", "multiplechoicegrid", "gridpilihanganda", "kisipilihanganda",
    },
    "checkbox_grid": {
        "checkboxgrid", "gridcheckbox", "gridkotakcentang",
    },
}


QUESTION_TYPES = {
    "short_text", "paragraph", "multiple_choice", "checkbox", "dropdown",
    "linear_scale", "date", "time", "grid", "checkbox_grid",
}


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def _display(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise GFormImportError("Encoding file teks tidak dapat dibaca.")


def _read_csv(data: bytes, filename: str) -> list[dict[str, Any]]:
    text = _decode_text(data)
    sample = text[:8192]
    delimiter = "\t" if Path(filename).suffix.lower() == ".tsv" else ","
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        pass
    rows = [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return [{"name": Path(filename).stem or "Data", "rows": rows, "state": "visible"}]


def _read_excel(data: bytes) -> list[dict[str, Any]]:
    if load_workbook is None:
        raise GFormImportError(
            "Dependensi openpyxl belum terpasang. Jalankan pip install -r requirements.txt."
        )
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise GFormImportError(f"Workbook tidak dapat dibaca: {exc}") from exc

    sheets: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_no > MAX_IMPORT_ROWS + 50:
                    raise GFormImportError(
                        f'Sheet "{sheet.title}" melebihi batas {MAX_IMPORT_ROWS:,} baris.'
                    )
                values = list(row[:MAX_IMPORT_COLUMNS])
                while values and values[-1] is None:
                    values.pop()
                rows.append(values)
            sheets.append(
                {"name": sheet.title, "rows": rows, "state": sheet.sheet_state}
            )
    finally:
        workbook.close()
    return sheets


def read_tables(data: bytes, filename: str) -> list[dict[str, Any]]:
    extension = Path(filename).suffix.lower()
    if extension in {".csv", ".tsv", ".txt"}:
        return _read_csv(data, filename)
    if extension in {".xlsx", ".xlsm"}:
        return _read_excel(data)
    if extension == ".xls":
        raise GFormImportError(
            "Format XLS lama belum didukung. Simpan ulang sebagai XLSX atau CSV."
        )
    raise GFormImportError(
        "Format tidak didukung. Gunakan XLSX, XLSM, CSV, TSV, atau TXT tabular."
    )


def _detect_header_row(rows: list[list[Any]]) -> int:
    best_index = 0
    best_score = float("-inf")
    all_aliases = set().union(*FIELD_ALIASES.values())

    for index, row in enumerate(rows[:30]):
        displayed = [_display(value) for value in row]
        nonempty = [value for value in displayed if value]
        if not nonempty:
            continue
        normalized = [_norm(value) for value in nonempty]
        alias_hits = sum(value in all_aliases for value in normalized)
        uniqueness = len(set(normalized)) / max(1, len(normalized))
        score = alias_hits * 14 + min(len(nonempty), 12) * 1.6 + uniqueness * 2
        score -= index * 0.08
        if len(nonempty) == 1 and alias_hits == 0:
            score -= 3
        if score > best_score:
            best_index = index
            best_score = score
    return best_index


def _columns_for_header(rows: list[list[Any]], header_index: int) -> list[dict[str, Any]]:
    width = max(
        [len(rows[header_index]) if header_index < len(rows) else 0]
        + [len(row) for row in rows[header_index + 1 : header_index + 12]]
        + [1]
    )
    width = min(width, MAX_IMPORT_COLUMNS)
    header = rows[header_index] if header_index < len(rows) else []
    seen: Counter[str] = Counter()
    columns = []
    for index in range(width):
        raw = _display(header[index]) if index < len(header) else ""
        label = raw or f"Kolom {index + 1}"
        seen[label] += 1
        if seen[label] > 1:
            label = f"{label} ({seen[label]})"
        columns.append({"index": index, "label": label, "raw": raw})
    return columns


def guess_mapping(columns: list[dict[str, Any]]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = {}
    for field, aliases in FIELD_ALIASES.items():
        match = next(
            (column["index"] for column in columns if _norm(column["raw"]) in aliases),
            None,
        )
        mapping[field] = match
    return mapping


def _sheet_quality(name: str, mapping: dict[str, int | None], row_count: int) -> int:
    normalized_name = _norm(name)
    score = sum(value is not None for value in mapping.values()) * 8
    if mapping.get("title") is not None:
        score += 20
    if mapping.get("row_type") is not None:
        score += 12
    if normalized_name in {"gformbuilder", "questions", "formbuilder", "items"}:
        score += 60
    if any(token in normalized_name for token in ("setting", "mapping", "skala", "score")):
        score -= 35
    score += min(row_count, 100) // 10
    return score


def _find_settings(tables: list[dict[str, Any]]) -> dict[str, str]:
    settings: dict[str, str] = {}
    for table in tables:
        if _norm(table["name"]) not in {"formsettings", "settings", "pengaturanform"}:
            continue
        for row in table["rows"][:100]:
            if len(row) < 2:
                continue
            key = _display(row[0])
            value = _display(row[1])
            if key and _norm(key) not in {"field", "key", "kunci"}:
                settings[key] = value
    return settings


def inspect_upload(data: bytes, filename: str) -> dict[str, Any]:
    tables = read_tables(data, filename)
    if not tables:
        raise GFormImportError("File tidak memiliki sheet atau baris yang dapat dibaca.")

    inspected = []
    for index, table in enumerate(tables):
        rows = table["rows"]
        header_index = _detect_header_row(rows)
        columns = _columns_for_header(rows, header_index)
        mapping = guess_mapping(columns)
        inspected.append(
            {
                "index": index,
                "name": table["name"],
                "state": table.get("state", "visible"),
                "row_count": max(0, len(rows) - header_index - 1),
                "column_count": len(columns),
                "header_row": header_index + 1,
                "columns": columns,
                "mapping": mapping,
                "head_rows": [
                    [_display(value) for value in row[:MAX_IMPORT_COLUMNS]]
                    for row in rows[:30]
                ],
                "sample_rows": [
                    [_display(value) for value in row[: len(columns)]]
                    for row in rows[header_index + 1 : header_index + 7]
                ],
                "quality": _sheet_quality(table["name"], mapping, len(rows)),
            }
        )

    selected = max(inspected, key=lambda item: item["quality"])["index"]
    settings = _find_settings(tables)
    title = (
        settings.get("Form_Title")
        or settings.get("Form Title")
        or Path(filename).stem.replace("_", " ")
    )
    description = settings.get("Form_Description") or settings.get("Form Description") or ""

    return {
        "filename": filename,
        "sheets": inspected,
        "selected_sheet": selected,
        "detected_title": title,
        "detected_description": description,
        "supported_types": [
            "section", "header", "short_text", "paragraph", "multiple_choice",
            "checkbox", "dropdown", "linear_scale", "date", "time", "grid",
            "checkbox_grid",
        ],
        "field_labels": FIELD_LABELS,
    }


def _cell(row: list[Any], mapping: dict[str, Any], field: str) -> str:
    index = mapping.get(field)
    if index is None or index == "":
        return ""
    try:
        index = int(index)
    except (TypeError, ValueError):
        return ""
    if index < 0 or index >= len(row):
        return ""
    return _display(row[index])


def _parse_bool(value: str, default: bool = False) -> bool:
    normalized = _norm(value)
    if normalized in {"true", "yes", "ya", "y", "1", "wajib", "required"}:
        return True
    if normalized in {"false", "no", "tidak", "n", "0", "opsional", "optional"}:
        return False
    return default


def _parse_options(value: str, separator: str = "auto") -> list[str]:
    value = str(value or "").strip()
    if not value:
        return []
    separators = {
        "pipe": r"\|",
        "semicolon": r";",
        "comma": r",",
        "newline": r"\r?\n",
    }
    if separator in separators:
        parts = re.split(separators[separator], value)
    elif "\n" in value or "\r" in value:
        parts = re.split(r"\r?\n", value)
    elif "|" in value:
        parts = value.split("|")
    elif ";" in value:
        parts = value.split(";")
    else:
        parts = [value]
    return [part.strip() for part in parts if part.strip()]


def _normalize_type(value: str) -> str:
    normalized = _norm(value)
    if not normalized:
        return ""
    for canonical, aliases in TYPE_ALIASES.items():
        if normalized == _norm(canonical) or normalized in aliases:
            return canonical
    return "unknown"


def _parse_int(value: str, default: int) -> int:
    try:
        return int(float(str(value).replace(",", ".")))
    except (TypeError, ValueError):
        return default


def normalize_upload(
    data: bytes,
    filename: str,
    *,
    sheet_index: int,
    header_row: int,
    mapping: dict[str, Any],
    form_options: dict[str, Any] | None = None,
) -> dict[str, Any]:
    options = form_options or {}
    tables = read_tables(data, filename)
    if sheet_index < 0 or sheet_index >= len(tables):
        raise GFormImportError("Sheet yang dipilih tidak ditemukan.")
    rows = tables[sheet_index]["rows"]
    header_index = max(0, header_row - 1)
    if header_index >= len(rows):
        raise GFormImportError("Nomor baris header berada di luar data.")

    data_rows = rows[header_index + 1 :]
    if len(data_rows) > MAX_IMPORT_ROWS:
        raise GFormImportError(f"Maksimum {MAX_IMPORT_ROWS:,} baris per impor.")

    default_required = bool(options.get("default_required", False))
    option_separator = str(options.get("option_separator", "auto"))
    warnings: list[str] = []
    items: list[dict[str, Any]] = []

    raw_types = [_normalize_type(_cell(row, mapping, "row_type")) for row in data_rows]
    has_explicit_sections = "section" in raw_types
    current_section = None

    for offset, row in enumerate(data_rows, start=header_index + 2):
        row_type_raw = _cell(row, mapping, "row_type")
        title = _cell(row, mapping, "title")
        description = _cell(row, mapping, "description")
        section = _cell(row, mapping, "section")
        section_description = _cell(row, mapping, "section_description")
        choices_raw = _cell(row, mapping, "options")
        row_type = _normalize_type(row_type_raw)

        mapped_values = [
            row_type_raw, title, description, section, section_description, choices_raw,
            _cell(row, mapping, "scale_min"), _cell(row, mapping, "scale_max"),
            _cell(row, mapping, "grid_rows"), _cell(row, mapping, "grid_columns"),
        ]
        if not any(mapped_values):
            continue

        if row_type == "unknown":
            warnings.append(
                f'Baris {offset}: tipe "{row_type_raw}" tidak dikenal dan dilewati.'
            )
            continue

        if not row_type:
            if not title and section:
                row_type = "section"
            elif _cell(row, mapping, "scale_min") or _cell(row, mapping, "scale_max"):
                row_type = "linear_scale"
            elif choices_raw:
                row_type = "multiple_choice"
            else:
                row_type = "short_text"

        if (
            not has_explicit_sections
            and row_type not in {"section", "header"}
            and section
            and section != current_section
        ):
            items.append(
                {
                    "type": "section",
                    "title": section,
                    "description": section_description,
                    "source_row": offset,
                }
            )
            current_section = section

        if row_type == "section":
            section_title = section or title
            if not section_title:
                warnings.append(f"Baris {offset}: bagian tanpa judul dilewati.")
                continue
            items.append(
                {
                    "type": "section",
                    "title": section_title,
                    "description": section_description or description,
                    "source_row": offset,
                }
            )
            current_section = section_title
            continue

        if row_type == "header":
            if not title:
                warnings.append(f"Baris {offset}: header tanpa judul dilewati.")
                continue
            items.append(
                {
                    "type": "header",
                    "title": title,
                    "description": description,
                    "source_row": offset,
                }
            )
            continue

        if not title:
            warnings.append(f"Baris {offset}: pertanyaan tanpa judul dilewati.")
            continue

        required_text = _cell(row, mapping, "required")
        item: dict[str, Any] = {
            "type": row_type,
            "title": title,
            "description": description,
            "required": _parse_bool(required_text, default_required),
            "source_row": offset,
        }

        if row_type in {"multiple_choice", "checkbox", "dropdown"}:
            choices = _parse_options(choices_raw, option_separator)
            if len(choices) < 2:
                warnings.append(
                    f"Baris {offset}: {row_type} membutuhkan minimal dua pilihan; dilewati."
                )
                continue
            item["choices"] = choices

        elif row_type == "linear_scale":
            minimum = _parse_int(_cell(row, mapping, "scale_min"), 1)
            maximum = _parse_int(_cell(row, mapping, "scale_max"), 5)
            if minimum not in {0, 1} or maximum < 3 or maximum > 10 or minimum >= maximum:
                warnings.append(
                    f"Baris {offset}: batas skala tidak valid; digunakan skala 1–5."
                )
                minimum, maximum = 1, 5
            item.update(
                {
                    "min": minimum,
                    "max": maximum,
                    "min_label": _cell(row, mapping, "scale_min_label"),
                    "max_label": _cell(row, mapping, "scale_max_label"),
                }
            )

        elif row_type in {"grid", "checkbox_grid"}:
            grid_rows = _parse_options(
                _cell(row, mapping, "grid_rows"), option_separator
            )
            grid_columns = _parse_options(
                _cell(row, mapping, "grid_columns") or choices_raw,
                option_separator,
            )
            if not grid_rows or len(grid_columns) < 2:
                warnings.append(
                    f"Baris {offset}: grid memerlukan baris dan minimal dua kolom; dilewati."
                )
                continue
            item.update({"rows": grid_rows, "columns": grid_columns})

        items.append(item)

    question_count = sum(item["type"] in QUESTION_TYPES for item in items)
    if question_count == 0:
        raise GFormImportError(
            "Tidak ada pertanyaan yang berhasil dikenali. Periksa pemetaan kolom dan tipe."
        )

    title = str(options.get("title") or Path(filename).stem.replace("_", " ")).strip()
    description = str(options.get("description") or "").strip()
    counts = Counter(item["type"] for item in items)

    return {
        "title": title,
        "description": description,
        "items": items,
        "settings": {
            "collect_email": bool(options.get("collect_email", False)),
            "progress_bar": bool(options.get("progress_bar", True)),
            "shuffle_questions": bool(options.get("shuffle_questions", False)),
            "create_response_sheet": bool(options.get("create_response_sheet", False)),
            "confirmation_message": str(options.get("confirmation_message") or "").strip(),
        },
        "summary": {
            "question_count": question_count,
            "item_count": len(items),
            "counts": dict(sorted(counts.items())),
            "source_sheet": tables[sheet_index]["name"],
        },
        "warnings": warnings,
    }


def public_preview(definition: dict[str, Any], limit: int = 120) -> dict[str, Any]:
    return {
        "title": definition["title"],
        "description": definition["description"],
        "summary": definition["summary"],
        "warnings": definition["warnings"],
        "items": definition["items"][:limit],
        "truncated": len(definition["items"]) > limit,
    }


SCRIPT_TEMPLATE = r'''/**
 * Generator Google Forms — dibuat oleh DocForge Universal GForm Builder.
 * Jalankan fungsi buatGoogleForm() satu kali dari editor Google Apps Script.
 */

const FORM_DEFINITION = __FORM_DEFINITION__;

function buatGoogleForm() {
  const definition = FORM_DEFINITION;
  const form = FormApp.create(definition.title, false);
  form.setDescription(definition.description || '');
  form.setIsQuiz(false);
  form.setCollectEmail(Boolean(definition.settings.collect_email));
  form.setProgressBar(Boolean(definition.settings.progress_bar));
  form.setShuffleQuestions(Boolean(definition.settings.shuffle_questions));
  if (definition.settings.confirmation_message) {
    form.setConfirmationMessage(definition.settings.confirmation_message);
  }

  definition.items.forEach(function (source) {
    let item;
    switch (source.type) {
      case 'section':
        item = form.addPageBreakItem().setTitle(source.title);
        setHelpText_(item, source.description);
        return;
      case 'header':
        item = form.addSectionHeaderItem().setTitle(source.title);
        setHelpText_(item, source.description);
        return;
      case 'short_text':
        item = form.addTextItem().setTitle(source.title);
        break;
      case 'paragraph':
        item = form.addParagraphTextItem().setTitle(source.title);
        break;
      case 'multiple_choice':
        item = form.addMultipleChoiceItem().setTitle(source.title);
        item.setChoiceValues(source.choices);
        break;
      case 'checkbox':
        item = form.addCheckboxItem().setTitle(source.title);
        item.setChoiceValues(source.choices);
        break;
      case 'dropdown':
        item = form.addListItem().setTitle(source.title);
        item.setChoiceValues(source.choices);
        break;
      case 'linear_scale':
        item = form.addScaleItem().setTitle(source.title);
        item.setBounds(source.min, source.max);
        item.setLabels(source.min_label || '', source.max_label || '');
        break;
      case 'date':
        item = form.addDateItem().setTitle(source.title);
        break;
      case 'time':
        item = form.addTimeItem().setTitle(source.title);
        break;
      case 'grid':
        item = form.addGridItem().setTitle(source.title);
        item.setRows(source.rows).setColumns(source.columns);
        break;
      case 'checkbox_grid':
        item = form.addCheckboxGridItem().setTitle(source.title);
        item.setRows(source.rows).setColumns(source.columns);
        break;
      default:
        throw new Error('Tipe tidak didukung: ' + source.type);
    }
    setHelpText_(item, source.description);
    item.setRequired(Boolean(source.required));
  });

  let responseSheetUrl = '';
  if (definition.settings.create_response_sheet) {
    const responseBook = SpreadsheetApp.create(definition.title + ' — Respons');
    form.setDestination(FormApp.DestinationType.SPREADSHEET, responseBook.getId());
    responseSheetUrl = responseBook.getUrl();
  }

  form.setPublished(true);
  const result = {
    editUrl: form.getEditUrl(),
    respondentUrl: form.getPublishedUrl(),
    responseSheetUrl: responseSheetUrl,
  };

  console.log('FORM SELESAI DIBUAT');
  console.log('Edit form: ' + result.editUrl);
  console.log('Form responden: ' + result.respondentUrl);
  if (result.responseSheetUrl) {
    console.log('Spreadsheet respons: ' + result.responseSheetUrl);
  }
  return result;
}

function setHelpText_(item, value) {
  const text = String(value || '').trim();
  if (text) item.setHelpText(text);
}
'''


def generate_apps_script(definition: dict[str, Any]) -> str:
    payload = {
        "title": definition["title"],
        "description": definition["description"],
        "items": definition["items"],
        "settings": definition["settings"],
    }
    serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    return SCRIPT_TEMPLATE.replace("__FORM_DEFINITION__", serialized)
